from docx import Document
from openpyxl import Workbook

# 打开Word文档
doc = Document('my.doc')

# 创建一个新的Excel工作簿
wb = Workbook()
ws = wb.active

# 逐行读取Word文档内容并写入Excel
for paragraph in doc.paragraphs:
    row_data = [paragraph.text]
    ws.append(row_data)

# 保存Excel文件
wb.save('converted_excel_file.xlsx')
